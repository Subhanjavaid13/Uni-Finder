import json
from datetime import date

from openpyxl import load_workbook

from scholar_radar.research_excel import build_workbook, fit_label, funding_realism, score_program

GOOD = {
    "country": "Italy", "region": "Calabria", "city": "Rende", "university": "Università della Calabria",
    "university_type": "public", "program_name": "MSc Artificial Intelligence and Computer Science",
    "field": "Artificial Intelligence", "program_url": "https://www.unical.it/ai", "language": "English",
    "application_fee_free": True, "moi_accepted": True, "ielts_required": False,
    "fully_funded_possible": "yes", "confidence": "high", "deadline_status": "last_year",
    "scholarships": [{"name": "Unical DSU", "covers": "tuition + stipend + housing",
                      "stipend": "EUR 5,000/year", "open_to_pakistan": True}],
    "documents": ["Transcripts", "CV"], "sources": ["https://www.unical.it/ai"],
}
WEAK = {
    "country": "Germany", "city": "Munich", "university": "Private Uni", "university_type": "private",
    "program_name": "MSc Data Science", "application_fee_free": "no", "ielts_required": "yes",
    "fully_funded_possible": "no", "confidence": "low",
}


def test_scoring():
    assert score_program(GOOD) >= 90 and fit_label(score_program(GOOD)) == "Top match"
    assert score_program(WEAK) < 50 and fit_label(score_program(WEAK)) == "Check"


def test_competitive_only_funding_is_downgraded():
    daad_only = {**GOOD, "scholarships": [
        {"name": "DAAD Study Scholarship - Pakistan", "type": "government", "need_based": False},
        {"name": "Deutschlandstipendium", "type": "government"}]}
    assert funding_realism(daad_only) == "competitive"
    assert score_program(daad_only) < score_program(GOOD)
    hungary = {**GOOD, "scholarships": [{"name": "Stipendium Hungaricum", "type": "government"}]}
    assert funding_realism(hungary) == "yes"
    elite_plus_discount = {**GOOD, "scholarships": [
        {"name": "Sparck AI Scholarship", "covers": "Full tuition + living allowance"},
        {"name": "ECS Global Scholarship", "type": "university", "covers": "GBP 5,000 tuition discount"}]}
    assert funding_realism(elite_plus_discount) == "competitive"
    turkish_private = {**GOOD, "scholarships": [
        {"name": "Sabanci Full Scholarship", "type": "university", "covers": "100% tuition + stipend + dorm"}]}
    assert funding_realism(turkish_private) == "yes"
    assert funding_realism({**GOOD, "fully_funded_possible": "partial"}) == "partial"


def test_build_workbook(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "italy.json").write_text(json.dumps({
        "region_name": "Italy South",
        "countries": [{"country": "Italy", "proof_of_funds": "EUR 6,000/year",
                       "national_scholarships": [{"name": "MAECI", "covers": "stipend"}]}],
        "scholarship_bodies": [{"country": "Italy", "region": "Calabria", "name": "Unical DSU",
                                "need_based": True, "open_to_pakistan": True, "url": "https://x.it"}],
        "programs": [GOOD, {**GOOD, "confidence": "low", "notes": "duplicate"}],
    }), encoding="utf-8")
    (raw / "germany.json").write_text(json.dumps({"programs": [WEAK]}), encoding="utf-8")
    (raw / "broken.json").write_text("{not json", encoding="utf-8")

    out = tmp_path / "out.xlsx"
    stats = build_workbook(raw, out, today=date(2026, 9, 15))
    assert stats == {"programs": 2, "top_matches": 1, "italy": 1, "scholarship_bodies": 1, "countries": 1}

    wb = load_workbook(out)
    assert wb.sheetnames == ["Read Me", "Top Matches", "Italy", "All Programmes",
                             "Universities by City", "Scholarships", "Visa & Country Info"]
    top = wb["Top Matches"]
    headers = [c.value for c in top[1]]
    row = {h: c for h, c in zip(headers, top[2])}
    assert row["University"].value == "Università della Calabria"
    assert row["Fit"].value == "Top match"
    assert row["MOI accepted?"].value == "Yes"
    assert row["Programme link"].hyperlink.target == "https://www.unical.it/ai"
    assert "duplicate" not in (row["Notes"].value or "")  # high-confidence duplicate kept
    assert wb["All Programmes"].max_row == 3
    assert wb["Scholarships"]["I2"].value == "Yes"


def test_key_findings_sheet_added_when_file_exists(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "a.json").write_text(json.dumps({"programs": [GOOD]}), encoding="utf-8")
    (tmp_path / "key_findings.json").write_text(json.dumps([
        {"priority": "2 HIGH", "topic": "B"}, {"priority": "1 URGENT", "topic": "A"}]), encoding="utf-8")
    out = tmp_path / "out.xlsx"
    build_workbook(raw, out, today=date(2026, 9, 15))
    ws = load_workbook(out)["Key Findings"]
    assert [ws["C2"].value, ws["C3"].value] == ["A", "B"]

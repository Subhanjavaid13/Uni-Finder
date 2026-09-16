"""Builds the Excel workbook from the country research files in research/raw/*.json."""

from __future__ import annotations

import json
import logging
import re
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FIT_FILLS = {
    "Top match": PatternFill("solid", fgColor="BBF7D0"),
    "Good": PatternFill("solid", fgColor="FEF08A"),
    "Check": PatternFill("solid", fgColor="E5E7EB"),
}
LINK_FONT = Font(color="1D4ED8", underline="single")
CONFIDENCE_RANK = {"high": 3, "medium": 2, "low": 1}
MAX_CELL = 32000

# (header, key or callable, column width)
PROGRAM_COLUMNS: list[tuple[str, Any, int]] = [
    ("Fit", "fit", 11),
    ("Score", "score", 7),
    ("Country", "country", 14),
    ("Region", "region", 16),
    ("City", "city", 14),
    ("University", "university", 32),
    ("Type", "university_type", 8),
    ("Programme", "program_name", 38),
    ("Field", "field", 18),
    ("Language", "language", 10),
    ("Duration", "duration", 10),
    ("Fully funded possible", "fully_funded_possible", 11),
    ("Scholarships (name - covers - stipend)", "scholarships_text", 50),
    ("Application fee", "application_fee", 14),
    ("Free to apply?", "application_fee_free", 9),
    ("Tuition (non-EU)", "tuition_non_eu", 22),
    ("IELTS required?", "ielts_required", 9),
    ("MOI accepted?", "moi_accepted", 9),
    ("English requirement", "english_requirement", 34),
    ("Entrance test", "entrance_test", 16),
    ("Academic requirements", "academic_requirements", 40),
    ("Min GPA", "min_gpa", 12),
    ("Other requirements", "other_requirements", 30),
    ("Documents", "documents_text", 40),
    ("Application portal", "application_portal", 20),
    ("Application window", "application_window", 24),
    ("Deadline", "deadline", 14),
    ("Deadline status", "deadline_status", 13),
    ("Intake", "intake", 14),
    ("Housing support", "housing_support", 28),
    ("Visa support", "visa_support", 36),
    ("Programme link", "program_url", 30),
    ("University link", "university_url", 26),
    ("Sources", "sources_text", 40),
    ("Confidence", "confidence", 10),
    ("Notes", "notes", 40),
]

SCHOLARSHIP_COLUMNS: list[tuple[str, str, int]] = [
    ("Country", "country", 14),
    ("Region", "region", 18),
    ("Scholarship / body", "name", 40),
    ("Type", "type", 12),
    ("Covers", "covers", 40),
    ("Stipend amount", "stipend_amount", 24),
    ("Housing", "housing", 10),
    ("Need-based?", "need_based", 9),
    ("Open to Pakistan?", "open_to_pakistan", 9),
    ("Income documents needed", "income_documents_needed", 36),
    ("Typical window", "typical_window", 22),
    ("Link", "url", 36),
    ("Notes", "notes", 40),
]

COUNTRY_COLUMNS: list[tuple[str, str, int]] = [
    ("Country", "country", 14),
    ("Student visa", "student_visa", 26),
    ("How to apply from Pakistan", "how_to_apply_from_pakistan", 44),
    ("Proof of funds", "proof_of_funds", 36),
    ("Scholarship letter counts as funds?", "scholarship_letter_counts_as_funds", 32),
    ("Visa fee", "visa_fee", 14),
    ("Tuition overview (non-EU)", "tuition_overview_non_eu", 34),
    ("Post-study work", "post_study_work", 30),
    ("Degree recognition steps", "degree_recognition_steps", 30),
    ("National scholarships", "national_scholarships_text", 50),
    ("Notes", "notes", 40),
    ("Sources", "sources_text", 40),
]


# ---- loading & normalising -------------------------------------------------------

def to_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        v = value.strip().lower()
        if v in ("true", "yes", "y"):
            return True
        if v in ("false", "no", "n"):
            return False
    return None


def yes_no(value: Any) -> str:
    b = to_bool(value)
    return "Yes" if b is True else "No" if b is False else "?"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(_text(v) for v in value if v not in (None, ""))
    return str(value).strip()


def load_research(raw_dir: Path) -> tuple[list[dict], list[dict], list[dict], list[str]]:
    """Returns (programs, scholarship_bodies, countries, region_names)."""
    programs: list[dict] = []
    bodies: list[dict] = []
    countries: list[dict] = []
    regions: list[str] = []
    for path in sorted(raw_dir.glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            log.warning("Skipping %s: %s", path.name, exc)
            continue
        regions.append(data.get("region_name") or path.stem)
        programs += [p for p in data.get("programs", []) if isinstance(p, dict)]
        bodies += [b for b in data.get("scholarship_bodies", []) if isinstance(b, dict)]
        countries += [c for c in data.get("countries", []) if isinstance(c, dict)]
    return programs, bodies, countries, regions


def _key(*parts: Any) -> str:
    return "|".join(re.sub(r"[^a-z0-9]+", " ", _text(p).lower()).strip() for p in parts)


def dedupe(rows: list[dict], key_fields: tuple[str, ...]) -> list[dict]:
    best: dict[str, dict] = {}
    for row in rows:
        k = _key(*(row.get(f) for f in key_fields))
        current = best.get(k)
        if current is None or (CONFIDENCE_RANK.get(_text(row.get("confidence")).lower(), 0)
                               > CONFIDENCE_RANK.get(_text(current.get("confidence")).lower(), 0)):
            best[k] = row
    return list(best.values())


COMPETITIVE_AWARDS = re.compile(
    r"daad|deutschlandstipendium|konrad zuse|eliza|secai|relai|eiffel|esop|chevening|fulbright"
    r"|excellence fellowship|sp[aä]rck|commonwealth|government of ireland|goi-ies"
    r"|danish government|erasmus mundus",
    re.IGNORECASE,
)
LIVING_COSTS = re.compile(
    r"stipend|living|allowance|accommodation|housing|dorm|board|lodging|cash|meal|salary"
    r"|assistantship|monthly",
    re.IGNORECASE,
)


def funding_realism(p: dict) -> str:
    """Downgrade "yes" to "competitive" when full funding relies only on elite merit awards.

    Need-based regional scholarships (e.g. Italian DSU) and quota schemes (e.g. Stipendium
    Hungaricum, Turkiye Burslari) keep "yes" because an eligible student has a realistic chance.
    """
    funded = _text(p.get("fully_funded_possible")).lower()
    if funded != "yes":
        return funded or "unknown"
    has_competitive = False
    for s in (p.get("scholarships") or []):
        if not isinstance(s, dict):
            continue
        if to_bool(s.get("need_based")) or _text(s.get("type")).lower() == "regional":
            return "yes"
        if COMPETITIVE_AWARDS.search(_text(s.get("name"))):
            has_competitive = True
        elif LIVING_COSTS.search(f"{_text(s.get('name'))} {_text(s.get('covers'))}"):
            return "yes"  # a realistic award that also pays living costs
    # Only elite awards (plus tuition-only discounts) fund living costs here.
    return "competitive" if has_competitive else "yes"


def score_program(p: dict) -> int:
    """How well a programme fits the student: full funding, free to apply, MOI/no IELTS."""
    score = 0
    funded = funding_realism(p)
    score += {"yes": 40, "competitive": 25, "partial": 15, "unknown": 5}.get(funded, 0)

    fee_free = to_bool(p.get("application_fee_free"))
    score += 15 if fee_free is True else 3 if fee_free is None else 0

    if to_bool(p.get("moi_accepted")) is True:
        score += 15
    elif to_bool(p.get("ielts_required")) is False:
        score += 12
    elif to_bool(p.get("ielts_required")) is None:
        score += 3

    if _text(p.get("university_type")).lower() == "public":
        score += 5
    if "english" in _text(p.get("language")).lower():
        score += 5
    score += {"high": 5, "medium": 2}.get(_text(p.get("confidence")).lower(), 0)
    if any(to_bool(s.get("open_to_pakistan")) for s in p.get("scholarships") or [] if isinstance(s, dict)):
        score += 5
    if _text(p.get("deadline_status")).lower() == "confirmed_2027":
        score += 5
    tuition = _text(p.get("tuition_non_eu")).lower()
    if re.search(r"\b(free|no tuition|0 eur|eur 0|tuition-free)\b", tuition):
        score += 5
    return min(score, 100)


def fit_label(score: int) -> str:
    return "Top match" if score >= 70 else "Good" if score >= 50 else "Check"


def prepare_program(p: dict) -> dict:
    row = dict(p)
    row["score"] = score_program(p)
    row["fit"] = fit_label(row["score"])
    row["fully_funded_possible"] = funding_realism(p)
    row["application_fee_free"] = yes_no(p.get("application_fee_free"))
    row["ielts_required"] = yes_no(p.get("ielts_required"))
    row["moi_accepted"] = yes_no(p.get("moi_accepted"))
    row["scholarships_text"] = "\n".join(
        " - ".join(x for x in (_text(s.get("name")), _text(s.get("covers")), _text(s.get("stipend"))) if x)
        + (" [Pakistan: " + yes_no(s.get("open_to_pakistan")) + "]")
        for s in (p.get("scholarships") or []) if isinstance(s, dict)
    )
    row["documents_text"] = _text(p.get("documents"))
    row["sources_text"] = "\n".join(_text(s) for s in (p.get("sources") or []))
    return row


def prepare_country(c: dict) -> dict:
    row = dict(c)
    row["national_scholarships_text"] = "\n".join(
        " - ".join(x for x in (_text(s.get("name")), _text(s.get("covers")), _text(s.get("stipend")),
                               _text(s.get("typical_window"))) if x)
        for s in (c.get("national_scholarships") or []) if isinstance(s, dict)
    )
    row["sources_text"] = "\n".join(_text(s) for s in (c.get("sources") or []))
    return row


def prepare_body(b: dict) -> dict:
    row = dict(b)
    row["need_based"] = yes_no(b.get("need_based"))
    row["open_to_pakistan"] = yes_no(b.get("open_to_pakistan"))
    return row


# ---- writing -----------------------------------------------------------------------

def _cell_value(value: Any) -> Any:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", _text(value))
    return text[:MAX_CELL]


def write_table(ws: Worksheet, columns: list[tuple[str, Any, int]], rows: list[dict]) -> None:
    for col, (header, _, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(rows, 2):
        for col, (header, key, _) in enumerate(columns, 1):
            value = _cell_value(row.get(key))
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, str) and value.startswith("http") and "\n" not in value and " " not in value:
                cell.hyperlink = value
                cell.font = LINK_FONT
            if key == "fit" and value in FIT_FILLS:
                cell.fill = FIT_FILLS[value]
                cell.font = Font(bold=True)

    ws.freeze_panes = "D2" if columns is PROGRAM_COLUMNS else "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _program_sort(row: dict) -> tuple:
    return (-row["score"], _text(row.get("country")), _text(row.get("city")), _text(row.get("university")))


def _location_sort(row: dict) -> tuple:
    return (_text(row.get("country")), _text(row.get("region")), _text(row.get("city")),
            _text(row.get("university")), -row["score"])


def write_readme(ws: Worksheet, today: date, programs: list[dict], bodies: list[dict],
                 countries: list[dict], regions: list[str]) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    lines: list[tuple[str, str]] = [
        ("Europe CS / AI Master's & Scholarships", ""),
        ("Generated", today.isoformat()),
        ("Student profile", "Pakistani BSCS graduate (2027), CGPA 3.1/4.0, English-medium degree (MOI), "
                            "no IELTS yet, 0 years work experience, needs full funding."),
        ("Programmes", str(len(programs))),
        ("Scholarship bodies", str(len(bodies))),
        ("Countries with visa info", str(len(countries))),
        ("Research regions", ", ".join(regions)),
        ("", ""),
        ("HOW TO USE", ""),
        ("0. Key Findings", "Urgent deadlines, what you are not eligible for, and English/visa traps. Read first."),
        ("1. Top Matches", "Best-fit programmes (score 60+): fully funded possible, free to apply, MOI / no IELTS."),
        ("2. Italy", "Every Italian programme found, sorted by region and city."),
        ("3. All Programmes", "Everything, sorted by country and city. Use the filter arrows in the header row."),
        ("4. Universities by City", "One row per university: how many programmes and the best score."),
        ("5. Scholarships", "Regional (e.g. Italian DSU), university, government and foundation scholarships."),
        ("6. Visa & Country Info", "Student visa, proof of funds, fees and post-study work per country."),
        ("", ""),
        ("COLUMNS", ""),
        ("Fit / Score", "Score out of 100: fully funded possible (40), free application (15), MOI accepted or "
                        "no IELTS (15), public university, English-taught, confidence, Pakistan-eligible "
                        "scholarship, confirmed 2027 dates. Top match >= 70, Good >= 50."),
        ("Fully funded possible", "yes = a realistic scholarship covers tuition AND living costs "
                                  "(e.g. need-based Italian DSU, Stipendium Hungaricum); "
                                  "competitive = full funding only through elite merit awards such as "
                                  "DAAD, Eiffel or ETH ESOP (possible but hard with CGPA 3.1); "
                                  "partial = only part of the costs; no; unknown."),
        ("Yes / No / ?", "'?' means the research could not confirm it - check the official page."),
        ("Deadline status", "confirmed_2027 = page covers the 2027/28 intake; last_year = based on the "
                            "2026/27 cycle, dates usually similar; unknown."),
        ("Confidence", "How reliable the row is (high = official pages checked)."),
        ("", ""),
        ("IMPORTANT", ""),
        ("Visa 'sponsorship'", "Universities don't sponsor student visas like employers. You get a student visa "
                               "with the admission (or pre-enrolment) letter plus proof of funds. A scholarship "
                               "award letter can usually serve as that proof - see 'Visa & Country Info'."),
        ("Italy DSU scholarships", "Italian regional DSU scholarships are need-based: you prove low family income "
                                   "with documents from Pakistan (ISEE parificato). Awarded students typically get "
                                   "a tuition waiver, a yearly cash stipend and sometimes housing/meals."),
        ("Accuracy", "Collected automatically from official websites on the date above by AI research agents. "
                     "Fees, deadlines and rules change every year - always confirm on the official link "
                     "before applying. Never pay an agent for a 'guaranteed' admission or scholarship."),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ca = ws.cell(row=r, column=1, value=a)
        cb = ws.cell(row=r, column=2, value=b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
        if r == 1:
            ca.font = Font(bold=True, size=16)
        elif a.isupper() and not b:
            ca.font = Font(bold=True, size=12, color="1D4ED8")
        elif a:
            ca.font = Font(bold=True)


def universities_by_city(programs: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for p in programs:
        k = _key(p.get("country"), p.get("city"), p.get("university"))
        g = grouped.setdefault(k, {
            "country": p.get("country"), "region": p.get("region"), "city": p.get("city"),
            "university": p.get("university"), "university_type": p.get("university_type"),
            "university_url": p.get("university_url"), "programs": [], "score": 0,
        })
        g["programs"].append(_text(p.get("program_name")))
        g["score"] = max(g["score"], p["score"])
    rows = []
    for g in grouped.values():
        g["program_count"] = len(g["programs"])
        g["programs_text"] = "\n".join(g["programs"])
        g["fit"] = fit_label(g["score"])
        rows.append(g)
    return sorted(rows, key=lambda g: (_text(g["country"]), _text(g["region"]), _text(g["city"]),
                                       _text(g["university"])))


FINDING_COLUMNS: list[tuple[str, str, int]] = [
    ("Priority", "priority", 14),
    ("Region", "region", 16),
    ("Topic", "topic", 36),
    ("Finding", "finding", 90),
    ("What to do", "action", 40),
]

UNIVERSITY_COLUMNS: list[tuple[str, str, int]] = [
    ("Country", "country", 14),
    ("Region", "region", 18),
    ("City", "city", 16),
    ("University", "university", 38),
    ("Type", "university_type", 9),
    ("Programmes found", "program_count", 11),
    ("Programme names", "programs_text", 50),
    ("Best score", "score", 9),
    ("Best fit", "fit", 11),
    ("University link", "university_url", 34),
]


def build_workbook(raw_dir: Path, out_path: Path, today: date | None = None,
                   top_threshold: int = 60) -> dict[str, int]:
    today = today or date.today()
    programs_raw, bodies_raw, countries_raw, regions = load_research(raw_dir)
    programs = [prepare_program(p) for p in dedupe(programs_raw, ("university", "program_name"))]
    bodies = [prepare_body(b) for b in dedupe(bodies_raw, ("country", "region", "name"))]
    countries = [prepare_country(c) for c in dedupe(countries_raw, ("country",))]

    wb = Workbook()
    write_readme(wb.active, today, programs, bodies, countries, regions)
    wb.active.title = "Read Me"

    findings_path = raw_dir.parent / "key_findings.json"
    if findings_path.exists():
        try:
            findings = json.loads(findings_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            log.warning("Ignoring %s: %s", findings_path.name, exc)
            findings = []
        if findings:
            write_table(wb.create_sheet("Key Findings"), FINDING_COLUMNS,
                        sorted(findings, key=lambda f: _text(f.get("priority"))))

    top = sorted((p for p in programs if p["score"] >= top_threshold), key=_program_sort)
    italy = sorted((p for p in programs if _text(p.get("country")).lower() == "italy"), key=_location_sort)
    sheets = [
        ("Top Matches", PROGRAM_COLUMNS, top),
        ("Italy", PROGRAM_COLUMNS, italy),
        ("All Programmes", PROGRAM_COLUMNS, sorted(programs, key=_location_sort)),
        ("Universities by City", UNIVERSITY_COLUMNS, universities_by_city(programs)),
        ("Scholarships", SCHOLARSHIP_COLUMNS,
         sorted(bodies, key=lambda b: (_text(b.get("country")), _text(b.get("region")), _text(b.get("name"))))),
        ("Visa & Country Info", COUNTRY_COLUMNS, sorted(countries, key=lambda c: _text(c.get("country")))),
    ]
    for title, columns, rows in sheets:
        write_table(wb.create_sheet(title), columns, rows)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"programs": len(programs), "top_matches": len(top), "italy": len(italy),
            "scholarship_bodies": len(bodies), "countries": len(countries)}
